import hashlib
import io
import json
import re
import xml.etree.ElementTree as ET
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
D = ROOT / "docs/v1"
A = ROOT / "artifacts/review-20260905"
tasks = json.loads((D / "tasks.json").read_text())["tasks"]
checks = [
    ["Python 测试", "通过", "392 passed", 0, "pytest.log", "本地pytest；不代表真实平台结果"],
    [
        "Web 单测（清理后）",
        "通过",
        "40 passed",
        0,
        "web-tests-after-cleanup.log",
        "保留全部有效测试",
    ],
    ["Studio 单测", "通过", "16 passed", 0, "web-tests.log", "软件测试替身"],
    [
        "TS 类型（清理后）",
        "通过",
        "Web/Studio/合同通过",
        0,
        "typecheck-after-cleanup.log",
        "编译类型检查",
    ],
    ["Web/Studio 构建", "通过", "两应用构建通过", 0, "build.log", "Studio 大包警告"],
    [
        "Companion 单测+lint",
        "通过",
        "41 tests / 14 suites；lint成功",
        0,
        "android.log",
        "主机测试，无障碍真机未验收",
    ],
    ["Ruff", "失败", "24 errors", 1, "ruff.log", "原始静态检查"],
    ["格式检查", "失败", "27 files", 1, "format.log", "未为审计做批量格式改写"],
    ["Pyright", "失败", "12 errors", 1, "pyright.log", "Edge/Outbox类型合同待修"],
    [
        "Web lint（清理后）",
        "失败",
        "4 errors",
        1,
        "web-lint-after-cleanup.log",
        "原5个，去掉1个无用测试变量",
    ],
    ["安全扫描", "无效/需重验", "退出0，grep报错", 0, "security.log", "不能视为所有扫描规则通过"],
    ["默认E2E", "环境污染", "4173复用其他项目", 1, "e2e.log", "不作为本项目失败数量；Studio未运行"],
    [
        "独立端口Web smoke",
        "部分失败",
        "16通过/4失败/8既有跳过",
        1,
        "e2e-isolated.log",
        "两处非精确文本定位器；不是已确认滚动故障",
    ],
    ["商品专用E2E", "通过", "1 passed", 0, "e2e-product.log", "API拦截用例，不是真实数据库到真机"],
    [
        "真机盘点",
        "通过",
        "OnePlus Android14；四App已安装",
        0,
        "device.json",
        "项目无障碍服务未启用",
    ],
    [
        "本地API只读访问",
        "通过",
        "OpenAPI 111 paths",
        0,
        "live-api-summary.json",
        "部署版本/外部来源未核实",
    ],
    ["四平台真实发布", "未执行", "0/4平台本轮完整验收", None, "", "本轮审核没有发布外部内容"],
]
(A / "verification.json").write_text(
    json.dumps(
        {
            "date": "2026-09-05",
            "checks": [
                dict(
                    zip(
                        ["check", "result", "detail", "exitCode", "log", "limitation"],
                        r,
                        strict=True,
                    )
                )
                for r in checks
            ],
        },
        ensure_ascii=False,
        indent=2,
    )
)
capabilities = [
    [
        "来源库连接",
        "未发现完整实现",
        "待确认用户现有库类型及只读入口",
        "V1-03~07",
        "source检索；01审核F10",
        "未验收",
    ],
    [
        "商品领域",
        "已有独立Product/Media及CRUD",
        "Product与Content下发断点、原子保存、历史修订",
        "V1-10/14",
        "db.py；OperationsView.vue",
        "软件基础已有",
    ],
    [
        "素材领域",
        "上传/去重/标签/分组已有",
        "来源同步、处理就绪、可视化选择",
        "V1-08/09",
        "media_store.py；迁移0010",
        "软件基础已有",
    ],
    [
        "内容修订",
        "已有ContentRevision与媒体关联",
        "三平台覆盖、格式校验、公众号HTML",
        "V1-11",
        "content_payload.py；迁移0011",
        "软件基础已有",
    ],
    [
        "账号设备",
        "记录绑定及API已有",
        "真实账号指纹/scope/无障碍权限",
        "V1-04/12",
        "mobile_service.py；device.json",
        "未真实闭环",
    ],
    [
        "手机投递",
        "下载hash校验与相册导出已有",
        "任务级授权/URI与选择项对应/大文件",
        "V1-17/18",
        "MediaDeliveryCoordinator.kt",
        "未真机验收",
    ],
    [
        "发布编排",
        "Plan/Target/Snapshot/Outbox/Intent已有",
        "APK提交保护、统一结果语义、多目标",
        "V1-13~22",
        "db.py；services.py",
        "未真实闭环",
    ],
    [
        "闲鱼",
        "字段预填和选图配方",
        "Web资源断点/选错图/最终提交/对账",
        "V1-18~22",
        "xianyu_publish.py",
        "未真实发布验收",
    ],
    [
        "小红书/抖音/公众号",
        "App已安装；内容基础可复用",
        "执行器与实际账号权限尚待实现/确认",
        "V1-23~26",
        "TargetLocatorRegistry.kt；device.json",
        "未真实发布验收",
    ],
    [
        "上线",
        "调试APK和部署脚本已有",
        "真实会话/release/备份恢复/无USB验证",
        "V1-28~30",
        "build.gradle.kts；session.ts",
        "未验收",
    ],
]
platforms = [
    [
        "闲鱼",
        "商品/多图/价格/运费/成色",
        "APK为当前基础；API需独立确认权限",
        "7.27.90",
        "预填代码已有，未最终提交",
        "身份/精确媒体/字段回读/permit/对账",
        "单次商品发布并核对列表/详情",
        "V1-18~22",
    ],
    [
        "小红书",
        "图文笔记",
        "先探测官方分享能力；必要时受控APK UI",
        "8.50.1",
        "仅安装，未接执行器",
        "分享成功不等于发布；标题正文/图序/账号",
        "笔记结果+审核状态；图文首验",
        "V1-23/24",
    ],
    [
        "抖音",
        "单视频作品",
        "有正式scope优先API；否则已授权UI路线",
        "39.6.0",
        "未核实video.create实际授权",
        "上传→创建→审核；提交后禁自动切渠道",
        "视频ID/作品核验/账号一致",
        "V1-23/25",
    ],
    [
        "微信公众号",
        "公众号图文文章",
        "先核实后台权限和当前官方文档",
        "微信8.0.76（非公众号能力证明）",
        "主体/权限未知；文档正文未取到",
        "草稿/发布/群发分开，不能默认个人微信能发",
        "文章ID/URL+状态；只有草稿不算done",
        "V1-23/26",
    ],
]
unknowns = [
    [
        "U01",
        "商品库在哪、主键和字段是什么",
        "来源管理员",
        "库类型/只读入口/10条脱敏样本",
        "V1-03/05/06",
        "未确认",
        "平台核心开发可继续",
    ],
    [
        "U02",
        "素材库原文件如何访问",
        "来源管理员",
        "20条样本/附件获取/大小/游标",
        "V1-03/05/06",
        "未确认",
        "不把缩略图当原件",
    ],
    [
        "U03",
        "四平台账号授权与能力",
        "账号管理员",
        "账号指纹/当前scope/到期/允许模式",
        "V1-12/23~26",
        "未确认",
        "只阻塞权限不明的平台",
    ],
    [
        "U04",
        "公众号主体、发布和草稿权限",
        "公众号管理员",
        "后台权限截图或脱敏接口结果",
        "V1-26",
        "未确认",
        "微信安装不等于公众号授权",
    ],
    [
        "U05",
        "真机绑定环境和无障碍服务",
        "设备管理员",
        "确认环境、启用项目服务、只读回执",
        "V1-04",
        "服务未启用",
        "ADB盘点已完成",
    ],
    [
        "U06",
        "实际商品/素材格式及规模",
        "业务负责人",
        "原始样本、容量、并发账号数量",
        "V1-08/23/28",
        "未确认",
        "先一设备/一账号串行",
    ],
    [
        "U07",
        "Release签名及生产部署归属",
        "部署管理员",
        "secret_ref/环境地址/备份位置",
        "V1-29",
        "未确认",
        "不复制明文密钥到计划",
    ],
]
report = (D / "01-项目审核.md").read_text()
findings = []
for m in re.finditer(
    r"^### (F\d+) / (P\d)：([^\n]+)\n(.*?)(?=\n### |\n## |\Z)", report, re.M | re.S
):
    num, prio, title, body = m.groups()
    findings.append(
        [num, prio, title, body.strip(), "静态证据+本轮运行；具体限制见正文", "待实施修复"]
    )
# Build reusable, printable tables.
wb = Workbook()
wb.remove(wb.active)
wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
navy = "17324D"
teal = "087F8C"
light = "EAF2F8"
gray = "526577"


def sheet(name, headers, rows, widths=None, height=58):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "C2" if len(headers) > 4 else "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for c in ws[1]:
        c.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=navy)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="微软雅黑", size=10, color="243B53")
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if c.row % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F0F5F9")
        ws.row_dimensions[row[0].row].height = height
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[__import__("openpyxl").utils.get_column_letter(i)].width = (
            widths or {}
        ).get(i, 24)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "第 &P 页 / 共 &N 页"
    ws.oddFooter.right.text = "2026-09-05 重审基线"
    return ws


summary = [
    ["审核结论", "基础可复用；一键来源库到四平台发布尚未完成验收", "详见01-项目审核.md"],
    ["八个业务里程碑", "本轮完整验收0/8", "不是代码完成率为0；按证据核验，不沿用旧40%"],
    ["计划任务数", 30, "由任务明细维护"],
    ["已验收任务数", "=COUNTIF('开发任务'!G2:G31,\"done\")", "只有全部验收满足才改done"],
    ["任务验收比例", "=B5/B4", "任务数口径，无工作量加权；不代表已有源码比例"],
    ["估算下限人日", "=SUM('开发任务'!E2:E31)", "一名跨端开发者，审核模型输出"],
    ["估算上限人日", "=SUM('开发任务'!F2:F31)", "平台审批/资料等待另计"],
    ["日历参考", "单人约9–17周（5有效开发日/周）", "规划估计非承诺；并行不能直接除模型数"],
    [
        "首要断点",
        "Product保存后误走Content下发；相册索引可能选错图",
        "先V1-10/14/18；保留正确的既有测试",
    ],
    ["当前真机", "OnePlus Android14，项目无障碍服务未启用", "四平台App已安装；无实际发布动作"],
    [
        "本轮清理",
        "2个孤立mock源码 + 1个测试无用变量；撤下旧计划/制表脚本",
        "9个文件有清理前备份；有效测试保留",
    ],
    ["入口", "docs/v1/README.md", "03任务卡供模型执行；tasks.json为结构化来源"],
    ["版本口径", "本地源码重审，2026-09-05", "不推断远端服务已同步或平台权限已获批"],
]
ws = sheet("总览", ["项目", "当前值", "说明"], summary, {1: 25, 2: 70, 3: 75}, 42)
ws["B6"].number_format = "0%"
# row checks: header=1; task count row4; done row5; ratio row6.
rows = []
for t in tasks:
    rows.append(
        [
            t["id"],
            t["title"],
            t["phase"],
            ", ".join(t["dependencies"]) or "无",
            t["estimate_low"],
            t["estimate_high"],
            t["status"],
            t["existing"],
            t["blocker"] or "依赖满足后可领取",
            t["owner"],
            t["acceptance"],
            t["evidence"],
        ]
    )
ws = sheet(
    "开发任务",
    [
        "ID",
        "任务",
        "阶段",
        "依赖",
        "人日下限",
        "人日上限",
        "状态",
        "已有基础（非验收）",
        "阻塞/启动条件",
        "负责人",
        "完成验收",
        "证据目录",
    ],
    rows,
    {1: 12, 2: 35, 3: 22, 4: 32, 5: 12, 6: 12, 7: 16, 8: 58, 9: 46, 10: 16, 11: 78, 12: 32},
    92,
)
dv = DataValidation(type="list", formula1='"todo,doing,blocked,blocked_hardware,review,done"')
dv.error = "请选择有效状态"
dv.errorTitle = "状态错误"
dv.showErrorMessage = True
ws.add_data_validation(dv)
dv.add("G2:G31")
for val, col in [("done", "D9EAD3"), ("blocked", "FCE4D6"), ("review", "FFF2CC")]:
    ws.conditional_formatting.add(
        "G2:G31", FormulaRule(formula=[f'G2="{val}"'], fill=PatternFill("solid", fgColor=col))
    )
rows = []
for t in tasks:
    rows.append(
        [
            t["id"],
            t["title"],
            "\n".join(t["files"]),
            "\n".join(f"{i}. {s}" for i, s in enumerate(t["steps"], 1)),
            "\n".join(t["verification_commands"]),
            t["acceptance"],
            t["evidence"],
        ]
    )
sheet(
    "实施步骤",
    ["ID", "任务", "修改范围", "按顺序实施", "验收命令/操作", "通过标准", "证据目录"],
    rows,
    {1: 12, 2: 36, 3: 70, 4: 115, 5: 80, 6: 75, 7: 32},
    205,
)
phase = defaultdict(lambda: [0, 0, []])
for t in tasks:
    phase[t["phase"]][0] += t["estimate_low"]
    phase[t["phase"]][1] += t["estimate_high"]
    phase[t["phase"]][2].append(t["id"])
sheet(
    "阶段安排",
    ["阶段", "任务范围", "下限人日", "上限人日", "交付门", "外部等待"],
    [
        [
            p,
            ", ".join(x[2]),
            x[0],
            x[1],
            {
                "M0 基线": "来源盘点、设备健康、工程门",
                "M1 来源与内容": "真实库同步+商品素材可编辑",
                "M2 发布核心": "冻结快照+互斥+APK收件/媒体授权",
                "M3 闲鱼": "正确预填→一次提交→结果对账",
                "M4 三内容平台": "小红书/抖音/公众号分别验收",
                "M5 完整V1": "四平台批次+恢复+release+8里程碑",
            }[p],
            "账号权限/库入口/平台审核按实际等待，未计入人日",
        ]
        for p, x in phase.items()
    ],
    {1: 24, 2: 70, 3: 14, 4: 14, 5: 58, 6: 62},
    66,
)
sheet(
    "当前能力",
    ["能力", "已有实现", "缺口", "任务", "代码证据", "本轮状态"],
    capabilities,
    {1: 22, 2: 50, 3: 72, 4: 22, 5: 48, 6: 25},
    72,
)
sheet(
    "审核问题",
    ["ID", "优先级", "问题", "证据与影响", "证据口径", "状态"],
    findings,
    {1: 12, 2: 12, 3: 55, 4: 130, 5: 45, 6: 24},
    190,
)
sheet(
    "平台矩阵",
    ["平台", "V1内容类型", "渠道方案", "当前安装版本", "现状", "实施关键点", "独立验收", "任务"],
    platforms,
    {1: 22, 2: 34, 3: 62, 4: 43, 5: 55, 6: 68, 7: 57, 8: 22},
    112,
)
accept = []
for line in (D / "02-实施方案.md").read_text().splitlines():
    if re.match(r"\| A\d+", line):
        accept.append([x.strip() for x in line.strip("|").split("|")] + ["未执行", ""])
sheet(
    "验收用例",
    ["ID", "输入或故障", "期望结果", "执行状态", "证据"],
    accept,
    {1: 12, 2: 75, 3: 90, 4: 20, 5: 45},
    60,
)
sheet(
    "验证记录",
    ["检查", "结果", "明细", "退出码", "日志", "限制"],
    checks,
    {1: 34, 2: 24, 3: 45, 4: 14, 5: 46, 6: 85},
    65,
)
sheet(
    "待确认事项",
    ["ID", "待确认", "提供人", "所需输入", "阻塞任务", "状态", "独立推进"],
    unknowns,
    {1: 12, 2: 50, 3: 20, 4: 68, 5: 25, 6: 24, 7: 50},
    65,
)
cleanup = json.loads((A / "cleanup-manifest.json").read_text())
sheet(
    "清理记录",
    ["原路径", "处理", "原SHA256", "恢复方式"],
    [
        [
            x["path"],
            "移除测试无用解构变量"
            if x["action"].startswith("remove unused")
            else "移除孤立mock或旧计划/制表产物",
            x["sha256"],
            "artifacts/review-20260905/cleanup-backup.zip 内原相对路径",
        ]
        for x in cleanup
    ],
    {1: 100, 2: 48, 3: 76, 4: 80},
    50,
)
for ws in wb:
    ws.sheet_properties.tabColor = teal if ws.title in ["总览", "开发任务", "实施步骤"] else navy
output = ROOT.parent / "LAMDA云控系统_V1重审开发计划_20260905.xlsx"
wb.save(output)
# Cache these four deterministic overview formulas for read-only spreadsheet previews.
ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
buffer = io.BytesIO()
with zipfile.ZipFile(output) as zin, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zout:
    for info in zin.infolist():
        data = zin.read(info.filename)
        if info.filename == "xl/worksheets/sheet1.xml":
            tree = ET.fromstring(data)  # noqa: S314 - input is the workbook generated above.
            done = sum(t["status"] == "done" for t in tasks)
            values = {
                "B5": done,
                "B6": done / len(tasks),
                "B7": sum(t["estimate_low"] for t in tasks),
                "B8": sum(t["estimate_high"] for t in tasks),
            }
            for cell in tree.findall(".//m:c", ns):
                if cell.attrib.get("r") in values:
                    v = cell.find("m:v", ns)
                    if v is None:
                        v = ET.SubElement(cell, "{" + ns["m"] + "}v")
                    v.text = str(values[cell.attrib["r"]])
            data = ET.tostring(tree, encoding="utf-8")
        zout.writestr(info, data)
output.write_bytes(buffer.getvalue())
cached = load_workbook(output, data_only=True)
assert [cached["总览"][x].value for x in ["B5", "B6", "B7", "B8"]] == [
    done,
    done / len(tasks),
    45,
    84,
]
# Validate actual file contents and cross-document consistency.
r = load_workbook(output, data_only=False)
assert len(r["开发任务"]["A"]) == 31
assert len(tasks) == 30 and len({t["id"] for t in tasks}) == 30
assert sum(t["estimate_low"] for t in tasks) == 45
assert sum(t["estimate_high"] for t in tasks) == 84
assert r["总览"]["B5"].value == "=COUNTIF('开发任务'!G2:G31,\"done\")"
assert r["总览"]["B6"].value == "=B5/B4"
assert len(accept) == 18 and len(findings) == 12
assert len(r.sheetnames) == 11
for ws in r:
    assert ws.max_row > 1 and ws.freeze_panes
    for row in ws:
        for c in row:
            assert not (isinstance(c.value, str) and c.value in ["#REF!", "#DIV/0!", "#VALUE!"])
(A / "workbook-validation.json").write_text(
    json.dumps(
        {
            "sheets": r.sheetnames,
            "tasks": 30,
            "findings": 12,
            "acceptanceCases": 18,
            "estimatePersonDays": [45, 84],
            "formulaReferences": "validated; recalculates on opening in spreadsheet application",
            "sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        },
        ensure_ascii=False,
        indent=2,
    )
)
print(output)
print(
    "Validated 11 sheets, 30 tasks, 12 findings, 18 acceptance cases; estimates 45–84 person-days."
)
